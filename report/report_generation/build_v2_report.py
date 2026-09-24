#!/usr/bin/env python3
"""Build the evidence-based V2 Cassandra client-centric consistency report."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "report"
GEN = REPORT / "report_generation"
FIG = GEN / "figures" / "png"
ASSET = GEN / "v2_assets"
SUMMARY = json.loads((GEN / "v2_analysis_summary.json").read_text())
OUT = REPORT / "draft" / "Cassandra_Client_Centric_Consistency_Report_V2.docx"

NAVY = "17324D"
BLUE = "1779BA"
DARK_BLUE = "1F4D78"
MUTED = "607386"
GRID = "B8C7D3"
LIGHT_BLUE = "EAF5FB"
LIGHT_GRAY = "F2F4F7"
PANEL = "F6F9FC"
GREEN = "2E8B68"
LIGHT_GREEN = "EAF7F1"
AMBER = "C47A14"
LIGHT_AMBER = "FFF4DD"
RED = "C43D4B"
LIGHT_RED = "FCECEF"
PURPLE = "7656A5"
WHITE = "FFFFFF"
BLACK = "243746"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, name="Calibri", size=None, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_table_geometry(table, widths: list[int]):
    assert sum(widths) == TABLE_WIDTH, (widths, sum(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    set_run_font(p.add_run(str(text)), size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths, *, header_fill=NAVY, font_size=8.7, alignments=None, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        shade(hdr.cells[i], header_fill)
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        if ri % 2:
            for c in cells:
                shade(c, PANEL)
        for i, value in enumerate(row):
            align = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], value, bold=(first_col_bold and i == 0), size=font_size, align=align)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), BLUE); rpr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; run.append(t)
    hyperlink.append(run); paragraph._p.append(hyperlink)


def add_field(paragraph, instruction, display=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep])
    if display:
        set_run_font(run, size=9, color=MUTED)
        run.add_text(display)
    run._r.append(end)


def add_page_number(paragraph):
    add_field(paragraph, "PAGE", "1")


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5); normal.font.color.rgb = rgb(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.20
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (11.5, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        s = styles[name]
        s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = rgb(color)
        s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        s = styles[name]
        s.font.name = "Calibri"; s.font.size = Pt(10.5); s.font.color.rgb = rgb(BLACK)
        s.paragraph_format.left_indent = Inches(0.5)
        s.paragraph_format.first_line_indent = Inches(-0.25)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.15
    cap = styles["Caption"]
    cap.font.name = "Calibri"; cap.font.size = Pt(9); cap.font.italic = True; cap.font.color.rgb = rgb(MUTED)
    cap.paragraph_format.space_before = Pt(3); cap.paragraph_format.space_after = Pt(9)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Evidence Tag" not in styles:
        s = styles.add_style("Evidence Tag", WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Calibri"; s.font.size = Pt(8); s.font.italic = True; s.font.color.rgb = rgb(MUTED)
        s.paragraph_format.space_before = Pt(2); s.paragraph_format.space_after = Pt(8)
    if "Callout" not in styles:
        s = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Calibri"; s.font.size = Pt(10.5); s.font.color.rgb = rgb(NAVY)
        s.paragraph_format.left_indent = Inches(0.2); s.paragraph_format.right_indent = Inches(0.2)
        s.paragraph_format.space_before = Pt(8); s.paragraph_format.space_after = Pt(8)
        s.paragraph_format.line_spacing = 1.15
    if "Code Inline" not in styles:
        s = styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
        s.font.name = "Menlo"; s.font.size = Pt(8.5); s.font.color.rgb = rgb(DARK_BLUE)


def set_paragraph_shading(paragraph, fill, border=None):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); ppr.append(shd)
    if border:
        pbdr = OxmlElement("w:pBdr")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "8"); el.set(qn("w:color"), border); el.set(qn("w:space"), "5"); pbdr.append(el)
        ppr.append(pbdr)


def add_callout(doc, label, text, fill=LIGHT_BLUE, border=BLUE):
    p = doc.add_paragraph(style="Callout")
    set_paragraph_shading(p, fill, border)
    set_run_font(p.add_run(label + "  "), size=10.5, color=border, bold=True)
    set_run_font(p.add_run(text), size=10.5, color=NAVY)
    return p


def add_body(doc, text, *, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), size=10.5, bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]), size=10.5)
    else:
        set_run_font(p.add_run(text), size=10.5)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.75)
    set_run_font(p.add_run(text), size=10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_run_font(p.add_run(text), size=10.5)
    return p


def add_picture(doc, path: Path, caption: str, alt: str, width=6.35, source=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", caption.split(".")[0])
    cp = doc.add_paragraph(caption, style="Caption")
    if source:
        ep = doc.add_paragraph(f"Source: {source}", style="Evidence Tag")
        ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return shape


def page_break(doc):
    doc.add_page_break()


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def pct(n, d, places=2):
    return f"{100*n/d:.{places}f}%" if d else "N/A"


def nfmt(n): return f"{int(n):,}"


def table_citation(doc, text):
    p = doc.add_paragraph(text, style="Evidence Tag")
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    return p


def outcome(setup):
    c = SUMMARY["overall"][setup]
    total = sum(c.values()); evaluable = c.get("violation", 0)+c.get("no_violation_observed", 0)
    return total, evaluable, c.get("violation", 0), c.get("no_violation_observed", 0), c.get("inconclusive", 0)


def section_header_footer(section):
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(0.78); section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1); section.right_margin = Inches(1)
    section.header_distance = Inches(0.35); section.footer_distance = Inches(0.35)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("CASSANDRA CLIENT-CENTRIC CONSISTENCY • DRAFT V2"), size=8, color=MUTED, bold=True)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run("NUS DSA508  |  "), size=8, color=MUTED)
    add_page_number(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(92); p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("DRAFT V2 • EVIDENCE-BASED REPORT"), size=11, color=AMBER, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run("Client-Centric Consistency\nin Apache Cassandra"), size=30, color=NAVY, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(34)
    set_run_font(p.add_run("Randomized coordinators versus Cassandra token-aware driver routing"), size=15, color=DARK_BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(28)
    set_run_font(p.add_run("Three-node Docker cluster • RF=3 • ONE / QUORUM / ALL\nNormal operation • node failure • 2|1 network partition"), size=11.5, color=MUTED)
    add_callout(doc, "Evidence snapshot", "57,168 structurally eligible main histories from 19 completed runs. The 13 randomized runs are full-profile evidence; the six token-aware runs are reported as exploratory because their saved profile is smoke.", LIGHT_BLUE, BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(32)
    set_run_font(p.add_run("NUS DSA508 – Scalable Distributed Computing"), size=12, color=NAVY, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Project group • Member details to be completed before submission"), size=10.5, color=MUTED)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(36)
    set_run_font(p.add_run("24 September 2026"), size=10.5, color=MUTED)
    page_break(doc)


def build():
    doc = Document()
    setup_styles(doc)
    section_header_footer(doc.sections[0])
    core = doc.core_properties
    core.title = "Client-Centric Consistency in Apache Cassandra – Draft V2"
    core.subject = "NUS DSA508 Scalable Distributed Computing project"
    core.author = "Project group"
    core.keywords = "Apache Cassandra, client-centric consistency, tunable consistency, Docker, distributed systems"
    settings = doc.settings._element
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    add_cover(doc)

    # Executive summary
    add_heading(doc, "Executive summary", 1)
    add_body(doc, "This project examined whether finite application histories exhibited read-your-writes (RYW), monotonic-reads (MR), monotonic-writes (MW), and writes-follow-reads (WFR) violations on a three-node Apache Cassandra 5.0.9 cluster. The workload varied write and read consistency levels independently across ONE, QUORUM, and ALL, and ran under normal operation, an established node failure, and a bilateral 2|1 internode network partition.")
    add_body(doc, "Two coordinator-selection designs were compared without pooling their result populations. The earlier design independently selected a random Cassandra coordinator for each operation. The later design supplied the partition routing key and allowed the Python driver's TokenAwarePolicy(DCAwareRoundRobinPolicy(local_dc=\"dc1\")) to create the query plan. In both designs, the receiving Cassandra node was the coordinator for that request; the coordinator was never a separate service.")
    rtot, reval, rv, rnv, rinc = outcome("previous_randomized")
    ttot, teval, tv, tnv, tinc = outcome("cassandra_token_aware")
    add_callout(doc, "Main finding", f"The full-profile randomized evidence contained {nfmt(rv)} violations among {nfmt(reval)} evaluable histories ({pct(rv,reval)}), while the exploratory token-aware evidence contained {nfmt(tv)} among {nfmt(teval)} ({pct(tv,teval,3)}). Every observed violation occurred during a network partition. This difference cannot be attributed to stronger consistency from token awareness: {pct(SUMMARY['route']['previous_randomized']['changed_coordinator'],rtot)} of randomized histories changed coordinator, compared with {pct(SUMMARY['route']['cassandra_token_aware']['changed_coordinator'],ttot)} of token-aware histories.", LIGHT_GREEN, GREEN)
    add_bullet(doc, f"Randomized full-profile evidence: {nfmt(rtot)} histories; {nfmt(rnv)} no violation observed, {nfmt(rv)} violation, and {nfmt(rinc)} inconclusive.")
    add_bullet(doc, f"Token-aware exploratory evidence: {nfmt(ttot)} histories; {nfmt(tnv)} no violation observed, {nfmt(tv)} violation, and {nfmt(tinc)} inconclusive.")
    add_bullet(doc, "Normal operation produced no violations. Node failure produced no violations but many unavailable histories. Network partitions produced all violations and the largest inconclusive populations.")
    add_bullet(doc, "Operation errors were classified as inconclusive rather than violations. Unavailable was the dominant error class, matching the replica-response requirements of ALL and minority-side QUORUM requests.")
    add_bullet(doc, "A no-violation outcome is evidence about one saved history, not proof that Cassandra guarantees the corresponding client-centric model.")
    add_callout(doc, "Evidence-status note", "The CSV exporter marks 19 completed runs as structurally eligible. The action plan separately requires smoke tests to be excluded from confirmatory claims. Because all six token-aware runs retain profile=smoke, this V2 report treats their measurements as exploratory and identifies the need for at least one verified full-profile token-aware run before final submission.", LIGHT_AMBER, AMBER)
    page_break(doc)

    # Contents
    add_heading(doc, "Contents", 1)
    sections = [
        ("1", "Objectives and research questions"), ("2", "Cassandra and client-centric consistency"),
        ("3", "Deployment architecture"), ("4", "Experimental designs"), ("5", "Predictions"),
        ("6", "Method"), ("7", "Evidence and analysis method"), ("8", "Results"),
        ("9", "Trace case studies"), ("10", "Cross-design discussion"), ("11", "Limitations and validity"),
        ("12", "Reproduction"), ("13", "Conclusion"), ("", "References and AI-use disclosure"),
        ("A", "Included and excluded runs"), ("B", "Evidence schema and file map"),
    ]
    for n, title in sections:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run((n + "  ") if n else ""), size=11, color=BLUE, bold=True)
        set_run_font(p.add_run(title), size=11, color=NAVY)
    add_callout(doc, "Reading convention", "Configuration labels are always write consistency/read consistency. Thus ONE/QUORUM means a ONE write followed by a QUORUM read.", PANEL, GRID)
    page_break(doc)

    # 1
    add_heading(doc, "1. Objectives and research questions", 1)
    add_body(doc, "The project asks how Cassandra's tunable consistency, failure topology, and client routing affect the consistency histories visible to an application. The application supplies data, a partition key, and a consistency level. It does not select a storage node in the token-aware design.")
    for item in [
        "RQ1. Under which write/read consistency combinations were RYW, MR, MW, or WFR counterexamples observed?",
        "RQ2. How did node failure and a 2|1 internode partition change evaluability and operation availability?",
        "RQ3. How did independent random coordinator selection differ from Cassandra-driver token-aware and data-center-aware routing?",
        "RQ4. Did measured observations agree with predictions derived from RF=3 replica acknowledgement and quorum intersection?",
        "RQ5. What can the recorded finite histories establish, and what remains outside the experiment's observational reach?",
    ]: add_bullet(doc, item)
    add_body(doc, "The unit of analysis was one ordered history over a run-scoped key. A history received exactly one primary verdict: violation, no violation observed, or inconclusive. Low-level exceptions were retained as explanatory evidence and did not create additional trial outcomes.")
    add_picture(doc, FIG/"06_client_centric_models.png", "Figure 1. Client-centric consistency models and the finite-history violation witnesses used by the experiment.", "Four cards defining RYW, MR, MW, and WFR with ordered operations and violation criteria.", source="Project workload definitions and src/checks.py")

    # 2
    add_heading(doc, "2. Cassandra and client-centric consistency", 1)
    add_heading(doc, "2.1 Replication and the request coordinator", 2)
    add_body(doc, "Apache Cassandra partitions data by token and replicates each partition according to the keyspace replication strategy. Any Cassandra node receiving a client request may coordinate it. The coordinator identifies replicas, forwards reads or mutations, and waits for the number of responses required by the selected consistency level [1]. The coordinator is therefore a per-request role, not a permanent leader or fourth service.")
    add_body(doc, "The lab keyspace used NetworkTopologyStrategy with replication factor 3 in dc1. Since the cluster also had three nodes, every node was a replica for each experiment key. Cassandra documentation recommends NetworkTopologyStrategy even for a single datacenter [1].")
    add_heading(doc, "2.2 Tunable consistency", 2)
    rows = [
        ("ONE", "1 replica", "Highest fault tolerance among tested levels; may expose a minority-side view."),
        ("QUORUM", "2 replicas", "Survives one unavailable replica when two replicas remain mutually reachable."),
        ("ALL", "3 replicas", "Requires every replica; unavailable during an established one-node failure or 2|1 cut."),
    ]
    add_table(doc, ["Consistency level", "RF=3 responses", "Experimental meaning"], rows, [1900,1900,5560], alignments=[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT], first_col_bold=True)
    table_citation(doc, "Table 1 source: Apache Cassandra tunable-consistency documentation [1, 2].")
    add_body(doc, "For RF=3, a write and read acknowledgement set must overlap when W + R > 3. QUORUM/QUORUM therefore intersects, while ONE/QUORUM sums to 3 and does not force intersection. This arithmetic supports visibility predictions for acknowledged operations, but it does not by itself prove every causal or session guarantee. Cassandra also sends writes to all replicas; the write consistency level controls how many responses the coordinator waits for [1].")
    add_heading(doc, "2.3 Conflict resolution and repair", 2)
    add_body(doc, "Cassandra resolves ordinary conflicting mutations by timestamp. The experiment supplied explicit increasing timestamps for main histories and used a decreasing-timestamp control to isolate this rule. Read repair was evaluated separately with BLOCKING and NONE tables. Hinted handoff was disabled, so post-fault cluster health could not be treated as proof that a previously missed mutation converged; convergence would require direct data evidence or repair [3, 4].")

    # 3
    page_break(doc)
    add_heading(doc, "3. Deployment architecture", 1)
    add_body(doc, "The common deployment consisted of three Cassandra services (n1, n2, n3), one client container, one private Docker network named lab, and three persistent Cassandra volumes. Host-side Python orchestration controlled Docker Compose and fault injection. CQL traffic used port 9042; Cassandra internode traffic used ports 7000 and 7001.")
    add_picture(doc, FIG/"01_windows_deployment_architecture.png", "Figure 2. Windows deployment architecture.", "Nested Windows, WSL2, and Docker boundaries showing controller, client container, three Cassandra nodes, volumes, traffic types, and evidence output.", source="compose.yaml and experiment orchestration source")
    add_body(doc, "On Windows, Docker Desktop can use WSL2 as the Linux container environment. The host controller invokes Compose and writes evidence through the mounted repository. The client process that issues CQL still executes inside the client container.")
    add_picture(doc, FIG/"02_macos_deployment_architecture.png", "Figure 3. macOS deployment architecture.", "macOS host with experiment controller and mounted workspace surrounding Docker Desktop's managed Linux VM, client container, and Cassandra cluster.", source="compose.yaml and Docker Desktop architecture [6]")
    add_body(doc, "On macOS, Docker Desktop runs the Docker Engine and Linux containers inside a managed Linux VM. The command begins on macOS, while the CQL client and Cassandra processes execute inside Docker's Linux environment [6].")

    # 4
    page_break(doc)
    add_heading(doc, "4. Experimental designs", 1)
    add_heading(doc, "4.1 Design A: harness-controlled random coordinator", 2)
    add_body(doc, "The earlier design independently sampled an eligible Cassandra node for each operation. This deliberately increased route diversity and the chance that successive operations crossed a network partition. It is useful as a counterexample search but gives the experiment harness control that ordinary applications typically delegate to the driver.")
    add_heading(doc, "4.2 Design B: Cassandra driver policy", 2)
    add_body(doc, "The later design supplied the keyspace and encoded partition routing key on each SimpleStatement. TokenAwarePolicy ranked replicas and DCAwareRoundRobinPolicy(local_dc=\"dc1\") ordered local hosts. The application API could not name a Cassandra node. Retries, consistency downgrades, and speculative execution were disabled, and the actual coordinator and attempted hosts were recorded for each operation. DataStax describes token-aware routing as prioritizing replicas and the load-balancing policy as producing an ordered per-query host plan [5].")
    add_picture(doc, FIG/"03_routing_policy_comparison.png", "Figure 4. Coordinator-selection designs compared.", "Two panels comparing independent harness-selected coordinators with token-aware driver routing to an ordinary Cassandra coordinator node.", source="random-coordinator-v2 and cassandra-driver-policy-v3 implementations")
    add_picture(doc, FIG/"07_token_aware_driver_policy_detail.png", "Figure 5. Token-aware driver coordinator selection in detail.", "Application routing key flows through token calculation and replica-aware ranking to a query plan and Cassandra coordinator.", source="src/worker.py, src/routing.py, and DataStax driver documentation [5]")
    add_callout(doc, "Interpretation", "Routing changes which histories are sampled and which failure domains successive operations traverse. It does not change the acknowledgement rule for ONE, QUORUM, or ALL.", LIGHT_BLUE, BLUE)

    # 5
    page_break(doc)
    add_heading(doc, "5. Predictions", 1)
    add_body(doc, "Predictions were registered before interpreting the saved outcomes. They assume one writer per key, no TTL or deletion, explicit increasing timestamps, retries disabled, and the finite-history oracles described in Section 6.")
    pred_rows = [
        ("ONE/ONE", "No", "May fail", "May regress", "Dependency visibility may fail"),
        ("ONE/QUORUM", "No", "May fail", "Quorum reads should resist regression", "No general causal guarantee"),
        ("ONE/ALL", "Yes", "Expected if operations complete", "Expected if reads complete", "No general causal guarantee"),
        ("QUORUM/ONE", "No", "May fail", "May regress", "No general causal guarantee"),
        ("QUORUM/QUORUM", "Yes", "Expected for acknowledged write/read", "Quorum reads should resist regression", "No general causal guarantee"),
        ("QUORUM/ALL", "Yes", "Expected if operations complete", "Expected if reads complete", "No general causal guarantee"),
        ("ALL/ONE", "Yes", "Expected if write completes", "Stronger visibility", "No general causal guarantee"),
        ("ALL/QUORUM", "Yes", "Expected if write completes", "Stronger visibility", "No general causal guarantee"),
        ("ALL/ALL", "Yes", "Expected while available", "Expected while available", "No completed witness expected"),
    ]
    add_table(doc, ["Write/read", "W+R>3", "RYW", "MR", "MW/WFR"], pred_rows, [1450,1150,2150,2200,2410], font_size=8.0, alignments=[WD_ALIGN_PARAGRAPH.CENTER]*2+[WD_ALIGN_PARAGRAPH.LEFT]*3, first_col_bold=True)
    table_citation(doc, "Table 2 source: registered predictions in report/predictions.md; mechanism supported by [1–3].")
    add_heading(doc, "5.1 Scenario predictions", 2)
    add_bullet(doc, "Normal: operations should usually complete; weak combinations may still admit a counterexample, but the experiment is a finite search.")
    add_bullet(doc, "Node failure: ONE and QUORUM may remain available on two survivors; ALL should be unavailable while one replica is down.")
    add_bullet(doc, "Network partition: ONE may complete on either side; QUORUM can complete only through the two-node side; ALL cannot complete until healing.")
    add_bullet(doc, "Random routing should create more cross-coordinator and cross-partition histories than the stable driver query plan, increasing counterexample exposure.")

    # 6
    page_break(doc)
    add_heading(doc, "6. Method", 1)
    add_heading(doc, "6.1 Factorial workload", 2)
    add_body(doc, "Each round combined three scenarios, nine write/read consistency pairs, and four client-centric models, for 108 cells. Repetitions were controlled by the saved plan. The pooled evidence spans runs with 1, 10, 50, or 100 rounds, so the report always states the actual denominator rather than assuming equal cell counts.")
    add_table(doc, ["Dimension", "Values", "Count"], [
        ("Write CL", "ONE, QUORUM, ALL", "3"), ("Read CL", "ONE, QUORUM, ALL", "3"),
        ("Configuration", "Cartesian product of write and read CL", "9"),
        ("Model", "RYW, MR, MW, WFR", "4"), ("Scenario", "normal, node_failure, network_partition", "3"),
        ("Cells per round", "9 × 4 × 3", "108"),
    ], [2300,5260,1800], alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER], first_col_bold=True)
    add_heading(doc, "6.2 Operation histories and oracles", 2)
    workload_rows = [
        ("RYW", "client-1 writes a=1; client-1 reads a", "Read is missing or a≠1"),
        ("MR", "producer writes a=1; client-1 reads twice", "Second successful read is older than first"),
        ("MW", "client-1 writes a=1 then b=1; observer reads", "b=1 visible while a≠1"),
        ("WFR", "producer writes a=1; dependent reads a and writes b=1; observer reads", "b=1 visible while a≠1, after dependency was observed"),
    ]
    add_table(doc, ["Model", "Ordered history", "Violation witness"], workload_rows, [1100,4860,3400], font_size=8.5, first_col_bold=True)
    add_body(doc, "If any required operation returned an error, the history was inconclusive. WFR was also inconclusive if the dependent client did not observe the prerequisite; MW was inconclusive if the final observation did not expose the successor. These coverage rules prevent absence of evidence from being counted as safety.")
    add_heading(doc, "6.3 Fault injection and recovery", 2)
    add_picture(doc, FIG/"04_fault_injection_architecture.png", "Figure 6. Fault-injection architecture and recovery gate.", "Normal, node-failure, and network-partition panels with availability expectations and recovery procedure.", source="experiments/faults.py and docs/execution-guide.md")
    add_body(doc, "Node failure used docker compose kill -s SIGKILL on a uniformly sampled victim. The runner recorded its identity, required both survivors to report the exact victim IP as DN in consecutive observations, executed the workload, restarted the same service with its volume, and required every observer to report three UN nodes before continuing.")
    add_body(doc, "Network partitioning installed bilateral DROP rules in a project-only LAB_FAULT chain for TCP ports 7000/7001 across a sampled 2|1 cut. Port 9042 remained reachable from the client. The current configuration specifies 15 seconds of stabilization, a 15-second hold, a 20-second post-recovery settle period, a 90-second failure-detection timeout, and a 600-second recovery timeout. Recovery removed only project rules and verified healthy membership.")
    add_heading(doc, "6.4 Controls", 2)
    add_body(doc, "Read-repair controls compared BLOCKING and NONE tables across overlapping quorum components. Timestamp controls issued an ALL write at T+100, then another at T+50, then read at ALL; returning the T+100 value confirmed timestamp ordering. Control results were not pooled with the 57,168 main histories.")

    # 7
    page_break(doc)
    add_heading(doc, "7. Evidence and analysis method", 1)
    add_picture(doc, FIG/"05_experiment_evidence_pipeline.png", "Figure 7. Evidence pipeline from parameterized execution to report tables.", "Configuration, Python runner, client, Cassandra, raw evidence, validation, CSV exports, and report layer.", source="report/report_generation/export_evidence.py and verify_exports.py")
    add_heading(doc, "7.1 Inclusion and evidence status", 2)
    add_body(doc, "The exporter found 33 runs matching the two permitted design/schema pairs. Fourteen incomplete runs were retained for audit and excluded because completion.json recorded completed=false. Nineteen completed runs passed structural checks for identity, expected counts, unique trial IDs, unique keys, and JSON record form. The export verifier reconciled 74,024 main rows plus 1,668 control rows and validated the saved hashes.")
    evidence_rows = [
        ("Random coordinator", "random-coordinator-v2", "v3", "13", "12,780", "Full", "Primary measured population"),
        ("Token-aware driver", "cassandra-driver-policy-v3", "v4", "6", "44,388", "Smoke", "Exploratory measured population"),
        ("Excluded incomplete", "approved families", "mixed", "14", "16,856 saved main rows", "Incomplete", "Audit only; excluded from totals"),
    ]
    add_table(doc, ["Population", "Design", "Schema", "Runs", "Main histories", "Profile", "Use"], evidence_rows, [1450,2080,650,650,1250,800,2480], font_size=7.7, alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT], first_col_bold=True)
    table_citation(doc, "Table 5 source: run_level_violation_matrix.csv and export verifier output.")
    add_callout(doc, "Profile caveat", "The exporter's included_in_analysis flag is a structural eligibility flag. It does not override the report action plan's instruction to exclude smoke tests from confirmatory claims. Token-aware measurements remain useful for diagnosis and design comparison, but the final submission should add a verified full-profile token-aware run.", LIGHT_AMBER, AMBER)
    add_heading(doc, "7.2 Outcome vocabulary and denominators", 2)
    outcome_rows = [
        ("Violation", "All required observations completed and contradicted the model oracle.", "Included in evaluable denominator"),
        ("No violation observed", "Required observations completed and did not contradict the oracle.", "Included in evaluable denominator"),
        ("Inconclusive", "An operation failed or an oracle precondition was not exposed.", "Excluded from violation-rate denominator"),
    ]
    add_table(doc, ["Outcome", "Meaning", "Rate treatment"], outcome_rows, [2000,4700,2660], first_col_bold=True)
    add_body(doc, "Violation rate is violations divided by evaluable histories. Evaluability rate is evaluable histories divided by all attempted main histories. Both are needed: a configuration can appear safe simply because the required operations rarely completed.")

    # 8 Results
    page_break(doc)
    add_heading(doc, "8. Results", 1)
    add_callout(doc, "Result in one sentence", "Observed violations were confined to network partitions and weak/non-intersecting visibility paths; failures with stronger replica requirements appeared primarily as inconclusive unavailability.", LIGHT_GREEN, GREEN)
    add_heading(doc, "8.1 Overall outcomes", 2)
    overall_rows=[]
    for setup,label,status in [("previous_randomized","Random coordinator","Primary/full"),("cassandra_token_aware","Token-aware driver","Exploratory/smoke")]:
        total,ev,v,nv,inc=outcome(setup)
        overall_rows.append((label,status,nfmt(total),nfmt(ev),nfmt(v),nfmt(inc),pct(v,ev,3),pct(ev,total)))
    add_table(doc, ["Setup", "Evidence status", "Attempted", "Evaluable", "Violations", "Inconclusive", "Violation/evaluable", "Evaluability"], overall_rows, [1450,1300,900,900,850,1000,1450,1510], font_size=7.8, alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT]+[WD_ALIGN_PARAGRAPH.CENTER]*6, first_col_bold=True)
    table_citation(doc, "Table 7 source: eligible main_trial rows in trial_level_evidence.csv.")
    add_picture(doc, ASSET/"08_outcome_composition.png", "Figure 8. Outcome composition by routing setup and scenario.", "Stacked bars showing no violation observed, inconclusive, and violation counts for each setup and scenario.", source="trial_level_evidence.csv; 57,168 eligible main histories")
    add_heading(doc, "8.2 Scenario effects", 2)
    scenario_rows=[]
    for setup,label in [("previous_randomized","Random"),("cassandra_token_aware","Token-aware*")]:
        for sc,slabel in [("normal","Normal"),("node_failure","Node failure"),("network_partition","Network partition")]:
            c=SUMMARY['scenario'][setup+'|'+sc]; total=sum(c.values()); v=c.get('violation',0); nv=c.get('no_violation_observed',0); inc=c.get('inconclusive',0); ev=v+nv
            assessment = "[x] Expected" if sc in ("normal","node_failure") else "[~] Plausible"
            scenario_rows.append((label,slabel,nfmt(total),nfmt(ev),nfmt(v),nfmt(inc),pct(v,ev,2),pct(ev,total),assessment))
    add_table(doc, ["Setup", "Scenario", "Attempted", "Evaluable", "Viol.", "Inconcl.", "Viol./eval.", "Evaluability", "Assessment"], scenario_rows, [1000,1300,850,850,650,850,1000,1000,1860], font_size=7.4, alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT]+[WD_ALIGN_PARAGRAPH.CENTER]*7)
    add_body(doc, "Normal operation produced zero violations in both populations. One randomized normal WFR history was inconclusive because its dependency was not observed. Node failure produced no consistency violation; instead, 2,260 randomized and 8,221 token-aware histories were inconclusive, overwhelmingly because the requested consistency level could not be satisfied. Network partitions produced all 200 observed violations and most operation errors.")

    add_heading(doc, "8.3 Model and consistency-level effects", 2)
    model_rows=[]
    for setup,label in [("previous_randomized","Random"),("cassandra_token_aware","Token-aware*")]:
        for model in ["RYW","MR","MW","WFR"]:
            c=SUMMARY['model'][setup+'|'+model]; v=c.get('violation',0); nv=c.get('no_violation_observed',0); inc=c.get('inconclusive',0); ev=v+nv
            model_rows.append((label,model,nfmt(ev),nfmt(v),nfmt(inc),pct(v,ev,3)))
    add_table(doc, ["Setup", "Model", "Evaluable", "Violations", "Inconclusive", "Violation/evaluable"], model_rows, [1500,950,1300,1300,1600,2710], font_size=8.2, alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER]+[WD_ALIGN_PARAGRAPH.CENTER]*4)
    add_body(doc, "In the full-profile randomized population, RYW produced 92 violations, MR 43, MW 46, and WFR 13. The token-aware exploratory population produced six RYW violations and none for MR, MW, or WFR. The absence of non-RYW token-aware witnesses coincided with very limited route changes and should not be interpreted as a stronger session guarantee.")
    add_picture(doc, ASSET/"10_network_partition_heatmap.png", "Figure 9. Network-partition violation rates among evaluable histories.", "Two heatmaps of write/read configurations by client-centric model, with violation counts and evaluable denominators.", source="trial_level_evidence.csv; network_partition only")
    add_body(doc, "The randomized heatmap concentrates violations in ONE/ONE and the non-intersecting ONE/QUORUM and QUORUM/ONE pairs. No violation was observed in cells containing ALL or in QUORUM/QUORUM. The token-aware exploratory population shows the same broad boundary, with six RYW witnesses in ONE/ONE, ONE/QUORUM, or QUORUM/ONE.")

    add_heading(doc, "8.4 Routing exposure", 2)
    rr=SUMMARY['route']['previous_randomized']; tr=SUMMARY['route']['cassandra_token_aware']
    rnet=SUMMARY['scenario_route']['previous_randomized|network_partition']; tnet=SUMMARY['scenario_route']['cassandra_token_aware|network_partition']
    routing_rows=[
        ("Random coordinator",nfmt(rr['trials']),nfmt(rr['changed_coordinator']),pct(rr['changed_coordinator'],rr['trials']),nfmt(rnet['crossed_partition_cut']),pct(rnet['crossed_partition_cut'],rnet['trials'])),
        ("Token-aware driver*",nfmt(tr['trials']),nfmt(tr['changed_coordinator']),pct(tr['changed_coordinator'],tr['trials']),nfmt(tnet['crossed_partition_cut']),pct(tnet['crossed_partition_cut'],tnet['trials'])),
    ]
    add_table(doc, ["Setup", "Histories", "Changed coordinator", "Changed %", "Crossed cut (network)", "Crossed %"], routing_rows, [1900,1200,1700,1200,1900,1460], font_size=8.2, alignments=[WD_ALIGN_PARAGRAPH.LEFT]+[WD_ALIGN_PARAGRAPH.CENTER]*5, first_col_bold=True)
    add_picture(doc, ASSET/"09_routing_exposure.png", "Figure 10. Coordinator and partition-side exposure.", "Side-by-side bars comparing coordinator changes, network cut crossings, and distinct coordinators per history.", source="trial_level_evidence.csv routing fields")
    add_body(doc, "The random harness changed coordinator in 10,298 of 12,780 histories (80.58%). The driver-policy population changed coordinator in only 344 of 44,388 histories (0.77%). During network partitions, 2,689 of 4,260 randomized histories crossed the cut (63.12%), compared with 92 of 14,796 token-aware histories (0.62%). This exposure gap is large enough to dominate any naive comparison of violation percentages.")

    add_heading(doc, "8.5 Inconclusive histories and availability", 2)
    add_picture(doc, ASSET/"11_inconclusive_causes.png", "Figure 11. Inconclusive reasons and operation-error classes.", "Stacked bars showing operation errors and oracle precondition failures, plus error-class totals.", source="trial_level_evidence.csv reason and exception fields")
    error_rows=[
        ("Random", "Unavailable", "5,272"), ("Random", "ReadTimeout", "15"), ("Random", "WriteTimeout", "11"),
        ("Token-aware*", "Unavailable", "18,032"), ("Token-aware*", "ReadTimeout", "58"), ("Token-aware*", "WriteTimeout", "34"), ("Token-aware*", "NoHostAvailable", "36"),
    ]
    add_table(doc, ["Setup", "Primary first-error class", "Histories"], error_rows, [2400,4660,2300], alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER])
    add_body(doc, "Unavailable dominated because Cassandra could immediately determine that the requested replica count was not alive from the selected coordinator's view. Timeouts were rarer and indicate that the coordinator began waiting but did not obtain enough responses before the deadline. NoHostAvailable appeared only in the exploratory driver runs, when no eligible host completed the request. These are availability outcomes, not consistency violations.")

    add_heading(doc, "8.6 Control experiments", 2)
    add_table(doc, ["Setup", "Read-repair no regression", "Read-repair regression", "Read-repair inconclusive", "Timestamp expected"], [
        ("Random", "67", "43", "140", "125"),
        ("Token-aware*", "70", "49", "703", "411"),
    ], [1800,1900,1750,1950,1960], font_size=8.2, alignments=[WD_ALIGN_PARAGRAPH.LEFT]+[WD_ALIGN_PARAGRAPH.CENTER]*4)
    add_body(doc, "All 536 timestamp controls were classified control_expected. Read-repair controls had substantial inconclusive populations; regression and no-regression counts are therefore mechanism checks rather than main-model rates. The controls are reported separately and do not alter the 57,168-history denominator.")

    # 9 traces
    page_break(doc)
    add_heading(doc, "9. Trace case studies", 1)
    add_heading(doc, "9.1 Token-aware RYW counterexample (exploratory)", 2)
    add_body(doc, "Trial 20260921T233512Z_000000000135283f:network_partition:18:ONE/ONE:RYW ran while n3 was isolated. Client-1 wrote at ONE through n2, then read at ONE through n3. Both operations completed, but the read returned the initialized state {a:0,b:0}. The saved coordinator sequence n2→n3 and crossed_partition_cut=true establish that the application moved between partition sides. This is a direct RYW violation witness for the finite history.")
    add_table(doc, ["Step", "Logical client", "Operation", "CL", "Coordinator", "Observed result"], [
        ("1", "client-1", "write a=1", "ONE", "n2", "ok"),
        ("2", "client-1", "read a,b", "ONE", "n3", "{a:0,b:0}"),
    ], [700,1500,1900,900,1300,3060], alignments=[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT])
    add_callout(doc, "Assessment: [x] Expected", "ONE/ONE does not force read/write intersection. The trace directly proves a cross-cut route and stale application observation; it does not reveal the internal replica response selected for each operation.", LIGHT_GREEN, GREEN)

    add_heading(doc, "9.2 Random-routing WFR counterexample", 2)
    add_body(doc, "Trial 20260919T100850Z_000000000135283b:network_partition:0:ONE/ONE:WFR established the dependency before testing it. The producer wrote a=1 through n2; the dependent client read a=1 through n2; the dependent write b=1 went through n3; and the observer read through n1, returning {a:0,b:1}. The dependent write was visible without the state that triggered it.")
    add_table(doc, ["Step", "Logical role", "Operation", "CL", "Coordinator", "Observed result"], [
        ("1", "producer-1", "write a=1", "ONE", "n2", "ok"),
        ("2", "dependent", "read a,b", "ONE", "n2", "{a:1,b:0}"),
        ("3", "dependent", "write b=1", "ONE", "n3", "ok"),
        ("4", "observer-1", "read a,b", "ONE", "n1", "{a:0,b:1}"),
    ], [700,1500,1900,900,1300,3060], font_size=8.3, alignments=[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT])
    add_callout(doc, "Assessment: [x] Expected", "The dependency was observed before b=1 was issued, and the final read exposed b without a. The four saved coordinators and cross-cut flag make this a strong WFR counterexample for the defined workload.", LIGHT_GREEN, GREEN)

    add_heading(doc, "9.3 Inconclusive availability trace", 2)
    add_body(doc, "Trial 20260921T080637Z_00000000000003e9:network_partition:0:ONE/ALL:MR ran with n3 isolated. A producer write at ONE succeeded through n3. Both later reads requested ALL through n3 and returned Unavailable with required_replicas=3 and alive_replicas=1. Since the reads produced no usable values, the MR oracle correctly returned inconclusive rather than violation.")
    add_callout(doc, "Assessment: [x] Expected", "ALL cannot be satisfied on either side of a 2|1 partition. The trace demonstrates the intended availability cost of a strong response requirement.", LIGHT_AMBER, AMBER)

    # 10. Let the discussion follow the trace section naturally. A forced break
    # can create a blank page when the final trace already fills its page.
    discussion_heading = add_heading(doc, "10. Cross-design discussion", 1)
    discussion_heading.paragraph_format.page_break_before = True
    add_heading(doc, "10.1 What the results support", 2)
    add_body(doc, "The results support three bounded conclusions. First, the normal cluster produced no counterexample in either saved population. Second, established failures converted stronger replica requirements primarily into unavailability rather than successful stale observations. Third, network partitions plus weak or non-intersecting write/read paths created the conditions for every observed counterexample.")
    add_heading(doc, "10.2 Why routing changed the observed rate", 2)
    add_body(doc, "The random design actively moved successive operations among Cassandra coordinators and across the 2|1 cut. The token-aware policy generated a stable first-live-local-replica ordering for most histories. Since every node was a replica at RF=3, token awareness did not narrow the replica set, but driver state and query-plan reuse strongly narrowed actual coordinator transitions. The two designs therefore sampled different histories even when they requested the same consistency levels.")
    add_body(doc, "The lower exploratory token-aware violation rate is consistent with lower cross-cut exposure. It is not evidence that TokenAwarePolicy adds RYW, MR, MW, or WFR guarantees. A fairer causal comparison would stratify or match histories by coordinator changes and partition-side crossings, then compare outcomes within those exposure classes.")
    add_heading(doc, "10.3 Agreement with predictions", 2)
    discussion_rows=[
        ("No normal-operation violations", "Generally expected", "[x] Expected", "Healthy RF=3 cluster and short sequential histories provided broad visibility."),
        ("No node-failure violations", "Plausible with errors", "[x] Expected", "Unavailable histories were not misclassified as violations."),
        ("Weak-path partition violations", "Predicted", "[x] Expected", "ONE/ONE and non-intersecting pairs can expose partition-side divergence."),
        ("No QQ or ALL-cell violations", "Predicted while evaluable", "[x] Expected", "Intersection/participation blocked completed witnesses; many histories were unavailable."),
        ("Much lower token-aware rate", "Exposure-sensitive", "[ ] Insufficient exposure", "Cross-cut histories were 0.62% versus 63.12%, and token runs were smoke-profile."),
    ]
    add_table(doc, ["Observation", "Prediction", "Assessment", "Comment"], discussion_rows, [1900,1600,1700,4160], font_size=8.0, first_col_bold=True)
    add_heading(doc, "10.4 Availability versus safety", 2)
    add_body(doc, "Safety and availability must be read together. ALL removed many histories from the evaluable denominator during faults. A zero violation count in such cells can mean that the system rejected the operation rather than returned a value satisfying the client-centric model. Conversely, ONE preserved more evaluable histories but allowed minority-side observations. The report therefore avoids a single pass/fail score.")

    # 11
    add_heading(doc, "11. Limitations and validity", 1)
    for item in [
        "Profile validity: all token-aware runs are saved as profile=smoke. Their 44,388 histories are exploratory until a full-profile run passes the experiment-specific verifier.",
        "Unequal populations: run counts, rounds, and selected cells differ between the two designs; pooled rates are descriptive rather than randomized causal estimates.",
        "Exposure confounding: routing policy substantially changed coordinator and partition-side exposure. This is part of the design effect but prevents interpreting raw violation-rate differences as a consistency guarantee.",
        "Single physical host: three containers share Docker Desktop resources, storage, clocks, and failure modes. A container SIGKILL is not identical to loss of an independent server.",
        "Single datacenter and RF=3: results do not generalize directly to multiple datacenters, larger replica sets, LOCAL_* levels, or different snitches and rack layouts.",
        "Finite counterexample search: no-violation-observed does not prove a universal session property; rare schedules may be missed.",
        "Internal response limits: coordinator and attempted-host evidence are recorded, but the exact responding replica set for every successful operation is not always available. Some causal explanations remain plausible rather than proven.",
        "Hints disabled: healthy membership after recovery does not prove divergent data converged. Direct post-recovery reads or repair evidence would be needed.",
        "Read repair and timestamp semantics: table settings and client-supplied timestamps shape observations and must accompany any reproduction.",
        "Environment drift: eligible runs span several Git revisions, Python versions, Docker Compose versions, and macOS/Linux hosts. The run-level CSV retains these values for audit.",
    ]: add_bullet(doc, item)
    add_callout(doc, "Required before final submission", "Run at least one full-profile token-aware experiment from the final code revision, validate it with scripts/verify_randomized.py, regenerate both CSVs, and rebuild this report. Replace the cover's group-member line and recheck every reported denominator.", LIGHT_RED, RED)

    # 12
    page_break(doc)
    add_heading(doc, "12. Reproduction", 1)
    add_heading(doc, "12.1 Software and deployment parameters", 2)
    add_table(doc, ["Parameter", "Recorded value"], [
        ("Database image", "consistency-lab-cassandra:5.0.9"),
        ("Python Cassandra driver", "cassandra-driver 3.29.2"),
        ("Services", "n1, n2, n3, client"),
        ("Cluster / datacenter / rack", "ConsistencyLab / dc1 / rack1"),
        ("Tokens per node", "16"),
        ("Keyspace", "lab; NetworkTopologyStrategy; dc1 RF=3"),
        ("Tables", "blocking (BLOCKING read repair); no_repair (NONE)"),
        ("Hints", "disabled"),
        ("Speculative execution", "disabled"),
        ("Network", "Docker network lab; CQL 9042; internode 7000/7001"),
    ], [2600,6760], first_col_bold=True)
    add_heading(doc, "12.2 Main commands", 2)
    commands=[
        "python3 -m venv .venv",
        ".venv/bin/python -m pip install -r requirements-report.txt",
        "python3 scripts/run_randomized.py --config config/cassandra_driver_experiments.json --seed 20260927 --plan-only",
        "python3 scripts/run_randomized.py --config config/cassandra_driver_experiments.json --seed 20260927",
        "python3 scripts/verify_randomized.py results/cassandra_policy_<run>",
        "python3 report/report_generation/export_evidence.py",
        "python3 report/report_generation/verify_exports.py",
    ]
    for i,cmd in enumerate(commands,1):
        p=doc.add_paragraph()
        p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.space_after=Pt(3)
        set_run_font(p.add_run(f"{i}. "),size=9,color=MUTED,bold=True)
        set_run_font(p.add_run(cmd),name="Menlo",size=8.2,color=DARK_BLUE)
    add_heading(doc, "12.3 Reproduction sequence", 2)
    for text in [
        "Start Docker Desktop or Docker Engine and wait until the daemon is ready.",
        "Review the registered predictions and final configuration before executing a measured run.",
        "Run the plan-only command and confirm expected matrix coverage.",
        "Execute with an explicit seed; do not run concurrent controllers.",
        "If interrupted, preserve the partial directory and use --recover; use --resume only when the main phase is complete and the runner accepts the checkpoint.",
        "Validate the exact run directory. Do not edit raw evidence to make a failed run pass.",
        "Regenerate the two report CSVs, run verify_exports.py, rebuild the V2 report, and inspect the rendered pages.",
    ]: add_number(doc,text)
    add_body(doc, "The complete operational procedure, troubleshooting guidance, schema barriers, fault proof, and recovery semantics are documented in README.md and docs/execution-guide.md. Raw evidence remains under results/<run_id>/ and is linked from both CSVs by source path and SHA-256 fields.")

    # 13
    add_heading(doc, "13. Conclusion", 1)
    add_body(doc, "The experiment demonstrated that Cassandra's operation-level consistency and an application's client-centric history are related but not interchangeable. In this RF=3 deployment, healthy execution produced no saved counterexample. Established node failure primarily reduced availability for operations requiring unreachable replicas. Network partitions created both high inconclusive rates and every observed RYW, MR, MW, and WFR violation.")
    add_body(doc, "The full-profile randomized coordinator design exposed 194 violations, concentrated in ONE/ONE and non-intersecting ONE/QUORUM or QUORUM/ONE paths. The exploratory token-aware design exposed six RYW violations under the same broad conditions. The large difference followed an equally large difference in route exposure: the random design changed coordinator and crossed the partition far more often.")
    add_body(doc, "The strongest defensible conclusion is therefore conditional: weak replica-response combinations can expose client-centric counterexamples when successive operations traverse divergent partition views; stronger combinations often trade those observations for unavailability. Token-aware routing reduced the sampled opportunities for those histories, but it did not supply a session guarantee. A final full-profile token-aware run is required before turning the exploratory comparison into a submission-level claim.")

    # References
    page_break(doc)
    add_heading(doc, "References", 1)
    refs=[
        ("[1]", "Apache Cassandra Documentation, “Dynamo: replication and tunable consistency.”", "https://cassandra.apache.org/doc/stable/cassandra/architecture/dynamo"),
        ("[2]", "Apache Cassandra Documentation, “Consistency levels.”", "https://cassandra.apache.org/doc/stable/cassandra/developing/cql/consistency.html"),
        ("[3]", "Apache Cassandra Documentation, “Read repair.”", "https://cassandra.apache.org/doc/stable/cassandra/managing/operating/read_repair.html"),
        ("[4]", "Apache Cassandra Documentation, “Hints.”", "https://cassandra.apache.org/doc/latest/cassandra/managing/operating/hints.html"),
        ("[5]", "DataStax Drivers Documentation, “Load balancing with Cassandra drivers.”", "https://docs.datastax.com/en/datastax-drivers/managing/balance-load.html"),
        ("[6]", "Docker Documentation, “Networking on Docker Desktop.”", "https://docs.docker.com/desktop/features/networking/"),
        ("[7]", "Project repository artifacts: compose.yaml, configuration files, experiment workloads, verifier, raw results, and report-generation CSV exports.", "https://github.com/whitetiger1399/scalableDistributed"),
    ]
    for num,title,url in refs:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.25); p.paragraph_format.space_after=Pt(7)
        set_run_font(p.add_run(num+" "),size=9.5,color=NAVY,bold=True); set_run_font(p.add_run(title+" "),size=9.5); add_hyperlink(p,url,url)
    add_heading(doc, "AI-use disclosure", 1)
    add_body(doc, "OpenAI Codex was used to review the draft and action plan, calculate aggregate summaries from the project-generated CSV files, create result visualizations, restructure the report, and generate the DOCX layout. The experiments and raw evidence were produced by the project software, not by AI. AI-generated interpretations were constrained to saved measurements and cited documentation. The project group remains responsible for checking the source data, technical claims, citations, group details, and final submission.")

    # Appendix A
    page_break(doc)
    add_heading(doc, "Appendix A. Run inventory", 1)
    add_heading(doc, "A.1 Included structurally eligible runs", 2)
    included=[]
    for r in SUMMARY['runs']:
        included.append(("Random" if r['setup']=='previous_randomized' else "Token-aware*",r['run_id'],r['configured_rounds'],r['saved_main_trials'],r['validation_status']))
    add_table(doc,["Setup","Run ID","Rounds","Main histories","Structural status"],included,[1450,4450,850,1350,1260],font_size=7.7,alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER])
    table_citation(doc, "* Token-aware runs are structurally eligible but saved as profile=smoke; reported exploratorily.")
    add_heading(doc, "A.2 Excluded runs", 2)
    excluded=[]
    for r in SUMMARY['excluded_runs']:
        excluded.append(("Random" if r['setup']=='previous_randomized' else "Token-aware",r['run_id'],r['run_completed'],r['validation_status'],r['validation_notes']))
    add_table(doc,["Setup","Run ID","Completed","Status","Reason"],excluded,[1400,4350,1000,1100,1510],font_size=7.7,alignments=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT])
    add_body(doc, "All 14 exclusions were caused by completion_false. Their partial records remain in the trial-level CSV for audit but were filtered by included_in_analysis=false from every result total in this report.")

    # Appendix B
    add_heading(doc, "Appendix B. Evidence schema and file map", 1)
    add_table(doc,["Artifact","Role in this report"],[
        ("trial_level_evidence.csv", "One row per main history or control, with provenance, verdict, routing, fault, timing, and ordered-operation fields."),
        ("run_level_violation_matrix.csv", "One row per run, including completion, validation, counts, environment, and complete cell matrix."),
        ("v2_analysis_summary.json", "Derived aggregate and trace selection used by the V2 builder; reproducible from the two CSVs."),
        ("v2_assets/08–11", "Result charts derived from the eligible main-history rows."),
        ("figures/png/01–07", "Deployment, routing, fault, evidence, and model diagrams."),
        ("results/<run_id>/", "Raw immutable JSON/JSONL evidence referenced by CSV source paths and hashes."),
    ],[3000,6360],first_col_bold=True)
    add_heading(doc, "B.1 Main analysis filters", 2)
    for text in [
        "included_in_analysis == true",
        "record_type == main_trial",
        "included_in_main_violation_matrix == true",
        "setup populations kept separate",
        "violation rate denominator = violation + no_violation_observed",
        "controls excluded from main denominators",
    ]: add_bullet(doc,text)
    add_heading(doc, "B.2 Suggested final-submission checks", 2)
    for text in [
        "Replace project-group metadata on the cover.",
        "Add and validate a full-profile token-aware run, then regenerate exports and charts.",
        "Confirm that the final code revision matches the source hashes recorded by the measured run.",
        "Re-run the exporter and verifier; independently sample at least one violation and one inconclusive history.",
        "Render the final DOCX/PDF and inspect every page for clipping, table splits, caption placement, and hyperlink correctness.",
    ]: add_bullet(doc,text)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
