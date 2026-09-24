#!/usr/bin/env python3
"""Build the V4/V5-only Cassandra client-centric consistency report."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_v2_report import (
    AMBER, BLUE, DARK_BLUE, GREEN, GRID, LIGHT_AMBER, LIGHT_BLUE, LIGHT_GREEN,
    LIGHT_RED, MUTED, NAVY, PANEL, RED, WHITE, add_body, add_bullet, add_callout,
    add_heading, add_page_number, add_picture, add_table, page_break, rgb,
    set_run_font, setup_styles, table_citation,
)

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "report"
GEN = REPORT / "report_generation"
FIG = GEN / "figures" / "png"
ASSET = GEN / "v5_assets"
SUMMARY = json.loads((GEN / "v5_analysis_summary.json").read_text())
OUT = REPORT / "draft" / "Cassandra_Client_Centric_Consistency_Report_V3.docx"
V4 = "cassandra-driver-policy-evidence-v4"
V5 = "cassandra-driver-policy-evidence-v5"


def nfmt(value):
    return f"{int(value):,}"


def pct(num, den, places=2):
    return f"{100 * num / den:.{places}f}%" if den else "N/A"


def metric(schema):
    return SUMMARY["by_schema"][schema]


def metric2(schema, dimension, value):
    return SUMMARY[dimension][f"{schema}|{value}"]


def set_header_footer(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.72)
    section.bottom_margin = Inches(.72)
    section.left_margin = Inches(.92)
    section.right_margin = Inches(.92)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.3)
    hp = section.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("CASSANDRA CLIENT-CENTRIC CONSISTENCY • REPORT V3"), size=8, color=MUTED, bold=True)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run("NUS DSA508  |  "), size=8, color=MUTED)
    add_page_number(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("REPORT V3 • V4/V5 DRIVER-POLICY EVIDENCE"), size=11, color=AMBER, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("Client-Centric Consistency\nin Apache Cassandra"), size=29, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(25)
    set_run_font(p.add_run("Three-node and five-node token-aware deployments"), size=15, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("RF=3 • ONE / QUORUM / ALL\nNormal operation • node failure • network partition"), size=11.5, color=MUTED)
    add_callout(doc, "Frozen evidence scope", "50,760 full-profile client histories from seven completed V4/V5 runs. Historical harness-randomized results, smoke runs, and incomplete runs are excluded from all reported aggregates.", LIGHT_BLUE, BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    set_run_font(p.add_run("NUS DSA508 – Scalable Distributed Computing"), size=12, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Project group • Member details to be completed before submission"), size=10, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    set_run_font(p.add_run("24 September 2026"), size=10, color=MUTED)
    page_break(doc)


def outcome_rows():
    rows = []
    for schema, label in ((V4, "V4 · three nodes"), (V5, "V5 · five nodes")):
        c = metric(schema)
        rows.append((label, nfmt(c["attempted"]), nfmt(c["evaluable"]), nfmt(c["violation"]),
                     nfmt(c["no_violation_observed"]), nfmt(c["inconclusive"]),
                     pct(c["evaluable"], c["attempted"]), pct(c["violation"], c["evaluable"], 3)))
    c = SUMMARY["overall"]
    rows.append(("Combined", nfmt(c["attempted"]), nfmt(c["evaluable"]), nfmt(c["violation"]),
                 nfmt(c["no_violation_observed"]), nfmt(c["inconclusive"]),
                 pct(c["evaluable"], c["attempted"]), pct(c["violation"], c["evaluable"], 3)))
    return rows


def add_command(doc, command):
    paragraph = add_callout(doc, "Command", command, PANEL, GRID)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def build():
    doc = Document()
    setup_styles(doc)
    set_header_footer(doc.sections[0])
    doc.core_properties.title = "Client-Centric Consistency in Apache Cassandra – Report V3"
    doc.core_properties.subject = "Three-node and five-node token-aware Cassandra experiments"
    doc.core_properties.author = "Project group"
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    add_cover(doc)

    add_heading(doc, "Executive summary", 1)
    add_body(doc, "This project investigated read-your-writes (RYW), monotonic-reads (MR), monotonic-writes (MW), and writes-follow-reads (WFR) consistency in Apache Cassandra. The experiment varied write and read consistency independently across ONE, QUORUM, and ALL under normal operation, one-node failure, and an established internode network partition.")
    add_body(doc, "Both reported deployments used Cassandra-driver token-aware and data-center-aware routing. V4 used three Cassandra nodes with RF=3 and a 1|2 partition. V5 expanded the cluster to five nodes while retaining RF=3, creating three replicas and two non-replicas per key and using a balanced 2|3 partition. The application supplied a routing key; the driver selected the coordinator from its query plan.")
    add_callout(doc, "Main result", "Across 50,760 completed full-profile histories, seven finite RYW counterexamples were observed: six in V4 and one in V5. All occurred during network partitions and crossed the recorded partition cut. No MR, MW, or WFR counterexample was observed. This is finite evidence, not proof that those models always hold.", LIGHT_GREEN, GREEN)
    add_table(doc, ["Deployment", "Attempted", "Evaluable", "Viol.", "No viol.", "Inconc.", "Evaluability", "Viol./eval."], outcome_rows(),
              [1450,1100,1050,650,1050,950,1350,1760], font_size=7.8, first_col_bold=True)
    add_bullet(doc, "Normal operation produced no violations and was essentially fully evaluable.")
    add_bullet(doc, "Node failure and network partition chiefly reduced availability; 99% or more of inconclusive histories were classified as Unavailable in both deployments.")
    add_bullet(doc, "V5 increased measured coordinator-change exposure from 0.78% to 0.99% and cross-cut exposure from 0.62% to 0.74%, but token awareness still strongly preferred replicas.")
    add_bullet(doc, "Cluster expansion alone did not create broad client-centric violations. Consistency level, completed operation path, partition placement, and route sequence remained decisive.")
    page_break(doc)

    add_heading(doc, "Contents", 1)
    for number, title in [
        ("1", "Research questions and scope"), ("2", "Cassandra consistency model"),
        ("3", "Deployment architectures"), ("4", "Experiment design and predictions"),
        ("5", "Execution, fault injection, and evidence"), ("6", "Results"),
        ("7", "Violation witnesses"), ("8", "Interpretation and validity"),
        ("9", "Reproduction"), ("10", "Conclusion"), ("", "References and AI-use disclosure"),
        ("A", "Run manifest"), ("B", "Evidence schema and artifact map")]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run((number + "  ") if number else ""), size=11, color=BLUE, bold=True)
        set_run_font(p.add_run(title), size=11, color=NAVY)
    add_callout(doc, "Notation", "Configuration labels are write consistency/read consistency. ONE/QUORUM therefore means a ONE write followed by a QUORUM read.", PANEL, GRID)
    page_break(doc)

    add_heading(doc, "1. Research questions and scope", 1)
    add_body(doc, "The study asks how Cassandra's tunable consistency, driver routing, replica placement, and failures affect histories visible to an application. The unit of analysis is one ordered client history over a run-scoped key.")
    for text in [
        "RQ1. Which write/read consistency combinations expose RYW, MR, MW, or WFR counterexamples?",
        "RQ2. How do node failure and network partition affect evaluability and successful stale observations?",
        "RQ3. How does expanding from three nodes/RF=3 to five nodes/RF=3 change routing and fault exposure?",
        "RQ4. Do the measured outcomes agree with quorum-intersection and availability predictions?",
        "RQ5. Which claims are supported by the saved finite evidence, and which remain outside its reach?",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "Result-population rule", "Only main_trial rows with included_in_analysis=true and run_profile=full are aggregated. V4/V5 smoke and incomplete runs remain visible in the audit CSV. Historical randomized-coordinator experiments are outside this report's result population.", LIGHT_AMBER, AMBER)
    add_picture(doc, FIG / "06_client_centric_models.png", "Figure 1. Client-centric models and finite-history violation witnesses.", "Four ordered-history definitions for RYW, MR, MW, and WFR.", source="Project workload definitions and src/checks.py")

    add_heading(doc, "2. Cassandra consistency model", 1)
    add_heading(doc, "2.1 Coordinator and replica roles", 2)
    add_body(doc, "A Cassandra node receiving a client request becomes its coordinator. It identifies the replicas for the partition key, forwards requests, and waits for the responses required by the requested consistency level. The coordinator is a per-request role rather than a leader. In V5, two of the five nodes are non-replicas for any specific RF=3 key, so coordinator and replica are distinct concepts.")
    add_heading(doc, "2.2 Tunable consistency at RF=3", 2)
    add_table(doc, ["Level", "Required responses", "Fault implication"], [
        ("ONE", "1 replica", "Can complete on either viable side containing a replica; no intersection guarantee."),
        ("QUORUM", "2 replicas", "Can tolerate one unavailable replica; must reach two replicas."),
        ("ALL", "3 replicas", "Requires every replica; becomes unavailable if any replica is unreachable."),
    ], [1350,2100,5910], first_col_bold=True)
    add_body(doc, "For RF=3, write and read acknowledgement sets necessarily intersect when W + R > 3. QUORUM/QUORUM intersects, while ONE/QUORUM equals 3 and does not force an overlap. Cassandra sends a mutation toward all replicas, but the selected write level controls how many acknowledgements the coordinator requires before success [1,2].")
    add_heading(doc, "2.3 What token-aware routing changes", 2)
    add_body(doc, "The Python driver computes the token from the routing key and prioritizes local replicas through TokenAwarePolicy(DCAwareRoundRobinPolicy(local_dc=\"dc1\")). Routing changes which node coordinates the operation and which paths are sampled. It does not strengthen ONE, QUORUM, or ALL. Retries, speculative execution, and consistency downgrades were disabled so the trace retained a clear operation outcome [3].")
    add_picture(doc, FIG / "07_token_aware_driver_policy_detail.png", "Figure 2. Token-aware coordinator selection.", "Routing key to token, replica ranking, query plan, selected coordinator, and replica acknowledgements.", source="src/routing.py, src/worker.py, and driver documentation [3]")

    page_break(doc)
    add_heading(doc, "3. Deployment architectures", 1)
    add_heading(doc, "3.1 V4: three nodes, RF=3", 2)
    add_body(doc, "The V4 deployment ran n1–n3, a client container, and the host-side controller in one Docker Compose project. Every node was a replica for every experiment key because cluster size equalled RF. The network-partition scenario isolated one node from the other two for internode ports 7000/7001 while preserving client CQL access on 9042.")
    add_picture(doc, FIG / "02_macos_deployment_architecture.png", "Figure 3. Three-node Docker deployment on macOS.", "Host controller, Docker Desktop VM, client container, three Cassandra nodes, volumes, and evidence output.", source="compose.yaml and experiment orchestration source")
    add_heading(doc, "3.2 V5: five nodes, RF=3", 2)
    add_body(doc, "The expanded deployment adds n4 and n5 but keeps RF=3. Each key therefore has exactly three replicas and two non-replicas. The V5 evidence records the replica set, selected coordinator, whether that coordinator was a replica, and attempted hosts. A seeded balanced 2|3 partition blocks every cross-group internode edge.")
    add_picture(doc, FIG / "08_five_node_expanded_architecture.png", "Figure 4. Expanded five-node/RF=3 architecture.", "Five-node Cassandra deployment with three replicas, two non-replicas, token-aware driver routing, balanced partition cut, and V5 evidence fields.", source="compose.expanded.yaml and cassandra_driver_expanded_experiments.json")
    add_table(doc, ["Property", "V4 · three nodes", "V5 · five nodes"], [
        ("Cassandra nodes", "n1–n3", "n1–n5"), ("Replication factor", "3", "3"),
        ("Replicas per key", "3 of 3", "3 of 5"), ("Non-replicas per key", "0", "2"),
        ("Partition topology", "1|2", "2|3"), ("Evidence schema", "driver-policy-evidence-v4", "driver-policy-evidence-v5"),
    ], [2350,3505,3505], first_col_bold=True)

    page_break(doc)
    add_heading(doc, "4. Experiment design and predictions", 1)
    add_heading(doc, "4.1 Factorial design", 2)
    add_body(doc, "Each round covers three scenarios, nine write/read consistency pairs, and four session models, producing 108 main cells. Each history receives one verdict: violation, no violation observed, or inconclusive. Required-operation failures and missing oracle preconditions are inconclusive.")
    add_table(doc, ["Dimension", "Values", "Count"], [
        ("Write CL", "ONE, QUORUM, ALL", "3"), ("Read CL", "ONE, QUORUM, ALL", "3"),
        ("Model", "RYW, MR, MW, WFR", "4"), ("Scenario", "normal, node_failure, network_partition", "3"),
        ("Cells per round", "3 × 3 × 4 × 3", "108"),
    ], [2300,5260,1800], first_col_bold=True)
    add_heading(doc, "4.2 Workload oracles", 2)
    add_table(doc, ["Model", "Ordered history", "Violation condition"], [
        ("RYW", "same client writes a=1, then reads a", "successful read omits the acknowledged write"),
        ("MR", "same client performs two ordered reads", "second successful read is older"),
        ("MW", "same client writes a=1 then b=1; observer reads", "b is visible while a is absent"),
        ("WFR", "dependent reads a, writes b; observer reads", "b is visible while observed dependency a is absent"),
    ], [1050,4740,3570], font_size=8.3, first_col_bold=True)
    add_heading(doc, "4.3 Predictions", 2)
    add_bullet(doc, "Normal: operations should usually complete; no counterexample is expected in these short initialized histories.")
    add_bullet(doc, "Node failure: ONE and QUORUM may remain available where enough replicas remain; ALL should become unavailable if a required replica is down.")
    add_bullet(doc, "Network partition: ONE can complete through one reachable replica; QUORUM requires two mutually reachable replicas; ALL cannot complete while replicas are split.")
    add_bullet(doc, "Weak paths such as ONE/ONE can expose an acknowledged write on one partition side followed by a read on the other side.")
    add_bullet(doc, "Five nodes create non-replica coordinators and more possible routes, but token-aware routing should still prefer replicas, so cluster size alone should not sharply increase violation exposure.")

    page_break(doc)
    add_heading(doc, "5. Execution, fault injection, and evidence", 1)
    add_heading(doc, "5.1 Node failure", 2)
    add_body(doc, "The controller samples a victim node and invokes docker compose kill -s SIGKILL. It waits until surviving observers repeatedly report the victim as DN, runs the workload, restarts the same service with its persistent volume, and gates the next episode on all configured observers reporting the full cluster as UN. V4 and V5 use the node list selected by the experiment configuration.")
    add_heading(doc, "5.2 Network partition", 2)
    add_body(doc, "The controller selects a seeded partition grouping and installs bilateral DROP rules in the project LAB_FAULT chain for Cassandra internode TCP ports 7000/7001. V4 uses a 1|2 cut. V5 uses a 2|3 cut and blocks all six cross-group node pairs in both directions. CQL port 9042 remains reachable so the driver can contact coordinators on either side. Recovery removes only project rules, waits for settlement, and verifies membership before continuing.")
    add_picture(doc, FIG / "04_fault_injection_architecture.png", "Figure 5. Fault injection and recovery gates.", "Normal, node-failure, and partition scenarios with consistency-level availability and recovery.", source="experiments/faults.py and experiments/network_partition.py")
    add_heading(doc, "5.3 Evidence pipeline", 2)
    add_picture(doc, FIG / "05_experiment_evidence_pipeline.png", "Figure 6. Evidence and report-generation pipeline.", "Configuration and workload execution through raw evidence, validation, normalized CSVs, matrices, charts, and report claims.", source="report/report_generation export and verification scripts")
    add_body(doc, "Raw JSON remains the primary evidence. The exporter normalizes every saved history into trial_level_evidence.csv and every run into run_level_violation_matrix.csv. Source hashes, run identity, route sequence, operation errors, fault metadata, and inclusion flags preserve traceability. verify_exports.py reconciles the CSVs and their manifest hashes.")

    page_break(doc)
    add_heading(doc, "6. Results", 1)
    add_heading(doc, "6.1 Evidence population", 2)
    add_table(doc, ["Deployment", "Runs", "Attempted", "Evaluable", "Viol.", "Inconc.", "Viol./eval."], [
        ("V4 · three nodes", "5", "44,280", "26,160", "6", "18,120", "0.023%"),
        ("V5 · five nodes", "2", "6,480", "4,437", "1", "2,043", "0.023%"),
        ("Combined", "7", "50,760", "30,597", "7", "20,163", "0.023%"),
    ], [1750,700,1300,1300,850,1300,2160], font_size=8.2, first_col_bold=True)
    add_picture(doc, ASSET / "12_v4_v5_outcome_composition.png", "Figure 7. Outcome composition by deployment and scenario.", "Normalized stacked bars showing no violation observed, inconclusive, and violation outcomes for V4 and V5.", source="v5_analysis_summary.json")
    add_heading(doc, "6.2 Scenario results", 2)
    scenario_rows = []
    for schema, label in ((V4, "V4"), (V5, "V5")):
        for scenario in ("normal", "node_failure", "network_partition"):
            c = metric2(schema, "by_schema_scenario", scenario)
            scenario_rows.append((label, scenario.replace("_", " "), nfmt(c["attempted"]), nfmt(c["evaluable"]),
                                  nfmt(c["violation"]), nfmt(c["inconclusive"]), pct(c["violation"], c["evaluable"], 3)))
    add_table(doc, ["Version", "Scenario", "Attempted", "Evaluable", "Viol.", "Inconc.", "Viol./eval."], scenario_rows,
              [900,1900,1300,1300,850,1300,1810], font_size=8.0)
    add_body(doc, "All seven violations occurred during network partitions. Normal operation had no violations. Node failure produced no successful counterexample but made many histories unavailable. Network-partition evaluability was 32.8% in V4 and 38.6% in V5, reflecting the requested replica count and side reached by the operation.")

    add_heading(doc, "6.3 Results by client-centric model", 2)
    model_rows = []
    for schema, label in ((V4, "V4"), (V5, "V5")):
        for model in ("RYW", "MR", "MW", "WFR"):
            c = metric2(schema, "by_schema_model", model)
            model_rows.append((label, model, nfmt(c["attempted"]), nfmt(c["evaluable"]), nfmt(c["violation"]),
                               nfmt(c["inconclusive"]), pct(c["violation"], c["evaluable"], 3)))
    add_table(doc, ["Version", "Model", "Attempted", "Evaluable", "Viol.", "Inconc.", "Viol./eval."], model_rows,
              [900,1000,1450,1450,850,1400,2310], font_size=8.1)
    add_picture(doc, ASSET / "14_v4_v5_model_violation_rates.png", "Figure 8. Observed violation rate by client-centric model.", "V4 and V5 violations divided by evaluable histories for each model.", source="v5_analysis_summary.json")
    add_callout(doc, "Interpretation boundary", "Zero observed MR, MW, or WFR violations is not a guarantee. These oracles require a specific completed sequence: a regression, visible successor, or visible dependency followed by loss. Stable driver routing and operation unavailability limited those exposures.", LIGHT_AMBER, AMBER)
    add_heading(doc, "6.4 Routing exposure", 2)
    add_table(doc, ["Deployment", "Histories", "Changed coordinator", "Changed %", "Partition histories", "Crossed cut", "Crossed %"], [
        ("V4", "44,280", "344", "0.78%", "14,760", "92", "0.62%"),
        ("V5", "6,480", "64", "0.99%", "2,160", "16", "0.74%"),
    ], [1100,1200,1700,1100,1650,1100,1510], font_size=8.0, first_col_bold=True)
    add_picture(doc, ASSET / "13_v4_v5_routing_exposure.png", "Figure 9. Coordinator-change and partition-crossing exposure.", "Measured percentages for V4 and V5 full-profile histories.", source="trial_level_evidence.csv")
    add_body(doc, "V5 recorded 19,420 measured operations coordinated by a replica and 20 by a non-replica. This confirms that the five-node setup created a real distinction between coordinator and replica, while the token-aware policy continued to prioritize replicas. The modest rise in route diversity explains why expansion did not create a large increase in observed counterexamples.")

    page_break(doc)
    add_heading(doc, "7. Violation witnesses", 1)
    add_body(doc, "Each witness below contains successful required operations and a value sequence rejected by the model oracle. All witnesses are RYW histories under network partition and all changed coordinator across the recorded cut.")
    vrows = []
    for v in SUMMARY["violations"]:
        version = "V4" if v["evidence_schema"] == V4 else "V5"
        groups = v["partition_groups_json"] or "1|2 cut"
        vrows.append((version, v["run_id"][:15] + "…", v["config_write_read"], v["model"],
                      v["coordinator_node_sequence"], groups))
    add_table(doc, ["Version", "Run ID", "CL", "Model", "Coordinator path", "Partition groups"], vrows,
              [850,2350,1100,850,1700,2510], font_size=7.4)
    add_callout(doc, "V5 trace", "Run 20260924T025626Z_000000000135283f produced one ONE/ONE RYW witness. The coordinator changed from n1 to n4 across groups [n1,n3,n5] and [n2,n4], after which the read omitted the acknowledged write.", LIGHT_RED, RED)
    add_heading(doc, "7.1 Consistency-level pattern", 2)
    add_body(doc, "ONE/ONE produced three V4 witnesses and the sole V5 witness. V4 also produced one ONE/QUORUM and two QUORUM/ONE witnesses. No witness was observed for QUORUM/QUORUM or any cell involving ALL. In fault scenarios, stronger levels often became unavailable, which reduced the set of successful histories capable of exposing a stale value.")
    add_heading(doc, "7.2 Inconclusive outcomes", 2)
    add_table(doc, ["Deployment", "Reason", "Histories", "Share of inconclusive"], [
        ("V4", "Unavailable", "17,986", "99.26%"), ("V4", "ReadTimeout", "58", "0.32%"),
        ("V4", "NoHostAvailable", "36", "0.20%"), ("V4", "WriteTimeout", "34", "0.19%"),
        ("V4", "Oracle precondition absent", "6", "0.03%"), ("V5", "Unavailable", "2,038", "99.76%"),
        ("V5", "Oracle precondition absent", "5", "0.24%"),
    ], [1350,3500,1700,2810], font_size=8.2)
    add_body(doc, "Unavailable means the coordinator could determine that too few replicas were reachable for the requested consistency level. A timeout or unavailable response is an availability outcome, not a consistency counterexample, because the required completed observation is missing.")

    page_break(doc)
    add_heading(doc, "8. Interpretation and validity", 1)
    add_heading(doc, "8.1 What the evidence supports", 2)
    add_bullet(doc, "Weak completed paths can violate RYW during an established network partition.")
    add_bullet(doc, "Faults primarily reduced availability; inconclusive histories must remain outside the violation-rate denominator.")
    add_bullet(doc, "A five-node/RF=3 token-aware deployment can produce a cross-cut RYW witness while preserving Cassandra-native routing.")
    add_bullet(doc, "The V4 and V5 evaluable violation rates were both 0.023%, despite different cluster topology and sample size.")
    add_heading(doc, "8.2 What the evidence does not support", 2)
    add_bullet(doc, "Token-aware routing does not guarantee RYW, MR, MW, or WFR.")
    add_bullet(doc, "No observed MR/MW/WFR witness does not establish that those models always hold.")
    add_bullet(doc, "The data do not show that adding nodes directly weakens consistency. Expansion changes replica placement and possible schedules; observed histories still depend on routing, timing, and consistency level.")
    add_bullet(doc, "Counts from V4 and V5 are not a controlled causal estimate of cluster-size effect because the versions contain different numbers of runs and histories.")
    add_heading(doc, "8.3 Threats to validity", 2)
    add_bullet(doc, "All containers ran on one physical host, so CPU, storage, and Docker scheduling do not reproduce independent machines or wide-area latency.")
    add_bullet(doc, "The token-aware query plan often retained one coordinator across a short history, limiting cross-cut schedules.")
    add_bullet(doc, "The V5 population contains two completed full runs, so rare-event rates have substantial sampling uncertainty.")
    add_bullet(doc, "Network partitions targeted internode traffic while preserving CQL reachability. Other asymmetric or client-network failures were not tested.")
    add_bullet(doc, "The oracles detect concrete finite witnesses; they do not implement an exhaustive formal history checker for all possible causal histories.")
    add_callout(doc, "Comparative conclusion", "Moving from three to five nodes made replica versus non-replica routing observable and slightly increased route-change exposure. It did not materially increase the observed violation rate. The strongest practical change was a richer, more realistic topology and evidence schema.", LIGHT_BLUE, BLUE)

    page_break(doc)
    add_heading(doc, "9. Reproduction", 1)
    add_body(doc, "The deployment is selected through the experiment configuration. The existing three-node profile remains unchanged; the expanded profile selects compose.expanded.yaml, n1–n5, RF=3, and the V5 evidence schema.")
    add_heading(doc, "9.1 Three-node V4 run", 2)
    add_command(doc, "python3 scripts/run_randomized.py --config config/cassandra_driver_experiments.json")
    add_heading(doc, "9.2 Five-node V5 run", 2)
    add_command(doc, "python3 scripts/run_randomized.py --config config/cassandra_driver_expanded_experiments.json")
    add_heading(doc, "9.3 Export and validate", 2)
    for cmd in [
        "python3 report/report_generation/export_evidence.py",
        "python3 report/report_generation/verify_exports.py",
        ".venv/bin/python report/report_generation/build_v5_report_assets.py",
        "python3 report/report_generation/build_v3_report.py",
    ]:
        add_command(doc, cmd)
    add_body(doc, "The principal audit artifacts are trial_level_evidence.csv, run_level_violation_matrix.csv, export_manifest.json, v5_analysis_summary.json, and violation_matrix_v5.md. The CSV source-path and hash columns trace every row back to its run evidence.")

    add_heading(doc, "10. Conclusion", 1)
    add_body(doc, "The experiment observed seven RYW counterexamples in 50,760 full-profile Cassandra driver-policy histories. Every witness required a network partition, a coordinator change across the cut, and a completed operation sequence. Normal operation produced no counterexamples, while node failure and partition scenarios frequently produced unavailability rather than stale successful observations.")
    add_body(doc, "Expanding the cluster from three to five nodes while retaining RF=3 created two non-replica nodes per key and made token-aware replica preference measurable. The five-node deployment yielded one cross-cut RYW witness, but its aggregate violation rate matched V4 at 0.023% of evaluable histories. The results therefore support a narrower conclusion: topology expands the possible schedules, while the observed client-centric outcome depends on the actual route, partition side, consistency level, and availability of the required replicas.")

    page_break(doc)
    add_heading(doc, "References", 1)
    refs = [
        "[1] Apache Cassandra Documentation. Dynamo architecture: tunable consistency and request coordination.",
        "[2] Apache Cassandra Documentation. Data definition and NetworkTopologyStrategy.",
        "[3] DataStax Python Driver Documentation. TokenAwarePolicy and DCAwareRoundRobinPolicy.",
        "[4] Apache Cassandra Documentation. Read repair and hinted handoff.",
        "[5] Docker Documentation. Docker Desktop networking and Compose.",
        "[6] Terry, D. B. et al. Session Guarantees for Weakly Consistent Replicated Data. PDIS, 1994.",
    ]
    for ref in refs:
        add_body(doc, ref)
    add_heading(doc, "AI-use disclosure", 1)
    add_body(doc, "OpenAI Codex was used to assist with code review, evidence normalization, chart and diagram generation, report structure, wording, and formatting. Experimental outcomes were computed from repository evidence and retained with source paths and hashes. The project team remains responsible for verifying the implementation, executing the experiments, checking the cited documentation, and reviewing every submitted claim.")

    page_break(doc)
    add_heading(doc, "Appendix A. Run manifest", 1)
    run_rows = []
    with (GEN / "run_level_violation_matrix.csv").open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["evidence_schema"] not in {V4, V5}:
                continue
            run_rows.append(("V4" if row["evidence_schema"] == V4 else "V5", row["run_id"], row["profile"] or "missing",
                             row["run_completed"], row["included_in_analysis"], row["saved_main_trials"],
                             row["violation_count"], row["inconclusive_count"], row["validation_notes"] or "—"))
    add_table(doc, ["Ver.", "Run ID", "Profile", "Done", "Included", "Main", "Viol.", "Inconc.", "Note"], run_rows,
              [550,2420,850,650,800,850,650,850,1740], font_size=6.8)
    add_body(doc, "The report aggregate contains the five completed full V4 runs and two completed full V5 runs marked Included=true. Other rows remain for auditability and are excluded from every result denominator.")

    page_break(doc)
    add_heading(doc, "Appendix B. Evidence schema and artifact map", 1)
    add_table(doc, ["Artifact", "Role"], [
        ("results/cassandra_policy_<run_id>/", "Primary completion, environment, trial, control, and fault evidence."),
        ("trial_level_evidence.csv", "One normalized row per main/control record with ordered operation and routing fields."),
        ("run_level_violation_matrix.csv", "One row per V4/V5 run with outcome cells, inclusion checks, and hashes."),
        ("export_manifest.json", "Frozen row counts, approved schemas, and SHA-256 digests."),
        ("v5_analysis_summary.json", "Machine-readable full-profile V4/V5 aggregates and witness inventory."),
        ("violation_matrix_v5.md", "Academic tables, calculations, witness list, and interpretation."),
        ("figures/png and v5_assets", "Architecture diagrams and derived result charts used in this report."),
    ], [3100,6260], font_size=8.4, first_col_bold=True)
    add_heading(doc, "Evidence identifiers", 2)
    add_bullet(doc, "V4 design: cassandra-driver-policy-v3; schema: cassandra-driver-policy-evidence-v4.")
    add_bullet(doc, "V5 design: cassandra-driver-policy-expanded-v1; schema: cassandra-driver-policy-evidence-v5.")
    add_bullet(doc, "Routing: TokenAwarePolicy(DCAwareRoundRobinPolicy(local_dc=\"dc1\")).")
    add_bullet(doc, "V5 additions: cluster profile, configured nodes, partition groups, replica_nodes, selected_node, selected_is_replica, and attempted_nodes.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
